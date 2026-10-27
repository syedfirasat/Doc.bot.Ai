"""Local Doc.Bot server.

Serves the standalone Doc.Bot page and stores user-owned records in a local
Excel workbook. This is intentionally a localhost development server, not a
production multi-user backend. The browser keeps the OpenRouter key in local
storage; this server never receives or writes that key.
"""
from __future__ import annotations

import argparse
import json
import mimetypes
import os
import threading
import uuid
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, unquote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
DATA_DIR = BASE_DIR / "data"
DB_PATH = DATA_DIR / "docbot_users.xlsx"
MAX_BODY_BYTES = 6 * 1024 * 1024
LOCK = threading.RLock()

SHEETS: dict[str, list[str]] = {
    "Users": ["user_id", "first_seen_utc", "last_seen_utc"],
    "Reports": ["user_id", "id", "ts", "date", "title", "details"],
    "ImportedSheets": [
        "user_id",
        "id",
        "name",
        "source",
        "imported_at",
        "headers_json",
        "rows_json",
    ],
    "QueryLog": [
        "user_id",
        "id",
        "ts",
        "date",
        "query",
        "answers",
        "category",
        "urgency",
        "prescription",
        "reply",
        "rounds_json",
    ],
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean(value: Any, limit: int = 20_000) -> str:
    return str(value or "").replace("\x00", "").strip()[:limit]


def user_id_from(value: Any) -> str:
    # The browser generates this opaque ID. It is not an authentication token.
    return clean(value, 120) or "anonymous-local-user"


def json_safe(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if value is None:
        return ""
    return value


def ensure_sheet(ws, headers: list[str]) -> None:
    if ws.max_row == 1 and all(ws.cell(1, col).value is None for col in range(1, len(headers) + 1)):
        for col, header in enumerate(headers, start=1):
            ws.cell(1, col).value = header
    elif ws.max_row == 0:
        ws.append(headers)
    else:
        existing = [ws.cell(1, col).value for col in range(1, len(headers) + 1)]
        if existing != headers:
            # Existing workbooks from this project are expected to use the
            # same schema. Refuse to silently rewrite a user-created sheet.
            raise RuntimeError(f"Sheet {ws.title!r} has an unexpected header row")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(64 + min(len(headers), 26))}{max(ws.max_row, 1)}"
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, col).fill = PatternFill("solid", fgColor="0F9E8F")
        ws.cell(1, col).alignment = Alignment(horizontal="center")
        ws.column_dimensions[chr(64 + col)].width = min(max(len(header) + 3, 14), 28)


def open_or_create_workbook():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    if DB_PATH.exists():
        workbook = load_workbook(DB_PATH)
    else:
        workbook = Workbook()
        default = workbook.active
        default.title = "Users"
    for title, headers in SHEETS.items():
        if title not in workbook.sheetnames:
            workbook.create_sheet(title)
        ensure_sheet(workbook[title], headers)
    return workbook


def save_workbook(workbook) -> None:
    # Updating the workbook while Excel has it open can fail on Windows. The
    # caller gets a useful 503 instead of a silent data loss.
    workbook.save(DB_PATH)


def row_dicts(ws, headers: list[str], user_id: str) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {header: json_safe(values[index] if index < len(values) else "") for index, header in enumerate(headers)}
        if str(row.get("user_id", "")) == user_id:
            results.append(row)
    return results


def row_to_report(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": clean(row.get("id"), 150),
        "ts": int(row.get("ts") or 0) if str(row.get("ts") or "").isdigit() else row.get("ts") or 0,
        "date": clean(row.get("date"), 100),
        "title": clean(row.get("title"), 100),
        "details": clean(row.get("details"), 1_000),
    }


def row_to_import(row: dict[str, Any]) -> dict[str, Any]:
    try:
        headers = json.loads(row.get("headers_json") or "[]")
    except (TypeError, ValueError):
        headers = []
    try:
        values = json.loads(row.get("rows_json") or "[]")
    except (TypeError, ValueError):
        values = []
    return {
        "id": clean(row.get("id"), 150),
        "name": clean(row.get("name"), 120),
        "source": clean(row.get("source"), 120),
        "importedAt": clean(row.get("imported_at"), 100),
        "headers": headers if isinstance(headers, list) else [],
        "rows": values if isinstance(values, list) else [],
    }


def row_to_record(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": clean(row.get("id"), 150),
        "ts": int(row.get("ts") or 0) if str(row.get("ts") or "").isdigit() else row.get("ts") or 0,
        "date": clean(row.get("date"), 100),
        "query": clean(row.get("query"), 1_000),
        "answers": clean(row.get("answers"), 8_000),
        "category": clean(row.get("category"), 120),
        "urgency": clean(row.get("urgency"), 40),
        "prescription": clean(row.get("prescription"), 2_000),
        "reply": clean(row.get("reply"), 20_000),
        "rounds": clean(row.get("rounds_json"), 20_000),
    }


def read_user_data(workbook, user_id: str) -> dict[str, Any]:
    report_rows = row_dicts(workbook["Reports"], SHEETS["Reports"], user_id)
    import_rows = row_dicts(workbook["ImportedSheets"], SHEETS["ImportedSheets"], user_id)
    record_rows = row_dicts(workbook["QueryLog"], SHEETS["QueryLog"], user_id)
    reports = [row_to_report(row) for row in report_rows]
    imports = [row_to_import(row) for row in import_rows]
    records = [row_to_record(row) for row in record_rows]
    reports.sort(key=lambda item: str(item.get("ts", "")), reverse=True)
    imports.sort(key=lambda item: str(item.get("importedAt", "")), reverse=True)
    records.sort(key=lambda item: str(item.get("ts", "")), reverse=True)
    return {"reports": reports[:200], "imports": imports[:50], "records": records[:3_000]}


def touch_user(workbook, user_id: str) -> None:
    ws = workbook["Users"]
    now = utc_now()
    for row_number in range(2, ws.max_row + 1):
        if str(ws.cell(row_number, 1).value or "") == user_id:
            ws.cell(row_number, 3).value = now
            return
    ws.append([user_id, now, now])


def upsert_row(ws, headers: list[str], user_id: str, item_id: str, values: dict[str, Any]) -> None:
    user_col = headers.index("user_id") + 1
    id_col = headers.index("id") + 1
    target_row = None
    for row_number in range(2, ws.max_row + 1):
        if str(ws.cell(row_number, user_col).value or "") == user_id and str(ws.cell(row_number, id_col).value or "") == item_id:
            target_row = row_number
            break
    if target_row is None:
        target_row = ws.max_row + 1
    for col, header in enumerate(headers, start=1):
        ws.cell(target_row, col).value = values.get(header, "")


def remove_row(ws, user_id: str, item_id: str) -> int:
    removed = 0
    for row_number in range(ws.max_row, 1, -1):
        if str(ws.cell(row_number, 1).value or "") == user_id and str(ws.cell(row_number, 2).value or "") == item_id:
            ws.delete_rows(row_number, 1)
            removed += 1
    return removed


def make_id(prefix: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex}"


def save_report(workbook, user_id: str, item: dict[str, Any]) -> dict[str, Any]:
    item_id = clean(item.get("id"), 150) or make_id("rep")
    values = {
        "user_id": user_id,
        "id": item_id,
        "ts": item.get("ts") or int(datetime.now().timestamp() * 1000),
        "date": clean(item.get("date"), 100) or utc_now(),
        "title": clean(item.get("title"), 100),
        "details": clean(item.get("details"), 1_000),
    }
    upsert_row(workbook["Reports"], SHEETS["Reports"], user_id, item_id, values)
    return {"id": item_id, "ts": values["ts"], "date": values["date"], "title": values["title"], "details": values["details"]}


def save_import(workbook, user_id: str, item: dict[str, Any]) -> dict[str, Any]:
    item_id = clean(item.get("id"), 150) or make_id("sheet")
    headers = item.get("headers") if isinstance(item.get("headers"), list) else []
    rows = item.get("rows") if isinstance(item.get("rows"), list) else []
    values = {
        "user_id": user_id,
        "id": item_id,
        "name": clean(item.get("name"), 120),
        "source": clean(item.get("source"), 120),
        "imported_at": clean(item.get("importedAt"), 100) or utc_now(),
        "headers_json": json.dumps(headers[:20], ensure_ascii=False),
        "rows_json": json.dumps(rows[:120], ensure_ascii=False)[:250_000],
    }
    upsert_row(workbook["ImportedSheets"], SHEETS["ImportedSheets"], user_id, item_id, values)
    return {
        "id": item_id,
        "name": values["name"],
        "source": values["source"],
        "importedAt": values["imported_at"],
        "headers": headers[:20],
        "rows": rows[:120],
    }


def save_record(workbook, user_id: str, item: dict[str, Any]) -> dict[str, Any]:
    item_id = clean(item.get("id"), 150) or make_id("rec")
    values = {
        "user_id": user_id,
        "id": item_id,
        "ts": item.get("ts") or int(datetime.now().timestamp() * 1000),
        "date": clean(item.get("date"), 100) or utc_now(),
        "query": clean(item.get("query"), 1_000),
        "answers": clean(item.get("answers"), 8_000),
        "category": clean(item.get("category"), 120),
        "urgency": clean(item.get("urgency"), 40),
        "prescription": clean(item.get("prescription"), 2_000),
        "reply": clean(item.get("reply"), 20_000),
        "rounds_json": clean(item.get("rounds"), 20_000),
    }
    upsert_row(workbook["QueryLog"], SHEETS["QueryLog"], user_id, item_id, values)
    return row_to_record(values)


def body_json(handler: "DocBotHandler") -> dict[str, Any]:
    raw_length = int(handler.headers.get("Content-Length", "0") or 0)
    if raw_length > MAX_BODY_BYTES:
        raise ValueError("Request body is too large")
    raw = handler.rfile.read(raw_length) if raw_length else b"{}"
    value = json.loads(raw.decode("utf-8"))
    if not isinstance(value, dict):
        raise ValueError("JSON body must be an object")
    return value


class DocBotHandler(SimpleHTTPRequestHandler):
    """Static-file handler plus a tiny JSON API for the local Excel workbook."""

    server_version = "DocBotLocal/1.0"

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(STATIC_DIR), **kwargs)

    def log_message(self, fmt: str, *args: Any) -> None:
        # Keep terminal output useful while avoiding a line for every asset.
        if self.path.startswith("/api") or self.command in {"POST", "DELETE"}:
            super().log_message(fmt, *args)

    def send_json(self, value: Any, status: int = HTTPStatus.OK) -> None:
        payload = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("Access-Control-Allow-Origin", "*")
        self.end_headers()
        self.wfile.write(payload)

    def send_error_json(self, status: int, message: str) -> None:
        self.send_json({"ok": False, "error": message}, status)

    def send_workbook(self) -> None:
        with LOCK:
            workbook = open_or_create_workbook()
            save_workbook(workbook)
            content = DB_PATH.read_bytes()
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Disposition", 'attachment; filename="docbot_users.xlsx"')
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_OPTIONS(self) -> None:
        self.send_response(HTTPStatus.NO_CONTENT)
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, DELETE, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/api/health":
            with LOCK:
                workbook = open_or_create_workbook()
                counts = {title: max(workbook[title].max_row - 1, 0) for title in SHEETS}
                save_workbook(workbook)
            self.send_json({"ok": True, "database": str(DB_PATH), "counts": counts})
            return
        if parsed.path == "/api/download.xlsx":
            try:
                self.send_workbook()
            except Exception as exc:  # pragma: no cover - OS-specific file locks
                self.send_error_json(HTTPStatus.SERVICE_UNAVAILABLE, f"Could not read the Excel workbook: {exc}")
            return
        if parsed.path == "/api/bootstrap":
            query = parse_qs(parsed.query)
            user_id = user_id_from(query.get("userId", [""])[0])
            try:
                with LOCK:
                    workbook = open_or_create_workbook()
                    touch_user(workbook, user_id)
                    data = read_user_data(workbook, user_id)
                    save_workbook(workbook)
                self.send_json({"ok": True, "userId": user_id, **data})
            except Exception as exc:
                self.send_error_json(HTTPStatus.SERVICE_UNAVAILABLE, f"Excel storage is unavailable: {exc}")
            return
        if parsed.path.startswith("/api/"):
            self.send_error_json(HTTPStatus.NOT_FOUND, "Unknown API route")
            return
        if parsed.path == "/":
            self.path = "/index.html"
        super().do_GET()

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path not in {"/api/reports", "/api/imports", "/api/records"}:
            self.send_error_json(HTTPStatus.NOT_FOUND, "Unknown API route")
            return
        try:
            payload = body_json(self)
            user_id = user_id_from(payload.get("userId"))
            item = payload.get("item") if isinstance(payload.get("item"), dict) else payload
            with LOCK:
                workbook = open_or_create_workbook()
                touch_user(workbook, user_id)
                if parsed.path == "/api/reports":
                    saved = save_report(workbook, user_id, item)
                elif parsed.path == "/api/imports":
                    saved = save_import(workbook, user_id, item)
                else:
                    saved = save_record(workbook, user_id, item)
                save_workbook(workbook)
            self.send_json({"ok": True, "item": saved})
        except json.JSONDecodeError:
            self.send_error_json(HTTPStatus.BAD_REQUEST, "Invalid JSON")
        except ValueError as exc:
            self.send_error_json(HTTPStatus.BAD_REQUEST, str(exc))
        except Exception as exc:
            self.send_error_json(HTTPStatus.SERVICE_UNAVAILABLE, f"Excel storage is unavailable: {exc}")

    def do_DELETE(self) -> None:
        parsed = urlparse(self.path)
        query = parse_qs(parsed.query)
        user_id = user_id_from(query.get("userId", [""])[0])
        try:
            with LOCK:
                workbook = open_or_create_workbook()
                if parsed.path == "/api/records":
                    ws = workbook["QueryLog"]
                    removed = 0
                    for row_number in range(ws.max_row, 1, -1):
                        if str(ws.cell(row_number, 1).value or "") == user_id:
                            ws.delete_rows(row_number, 1)
                            removed += 1
                elif parsed.path in {"/api/reports", "/api/imports"} and parsed.path == "/api/reports":
                    removed = 0
                elif parsed.path.startswith("/api/reports/"):
                    item_id = unquote(parsed.path.rsplit("/", 1)[-1])
                    removed = remove_row(workbook["Reports"], user_id, item_id)
                elif parsed.path.startswith("/api/imports/"):
                    item_id = unquote(parsed.path.rsplit("/", 1)[-1])
                    removed = remove_row(workbook["ImportedSheets"], user_id, item_id)
                else:
                    self.send_error_json(HTTPStatus.NOT_FOUND, "Unknown API route")
                    return
                save_workbook(workbook)
            self.send_json({"ok": True, "removed": removed})
        except Exception as exc:
            self.send_error_json(HTTPStatus.SERVICE_UNAVAILABLE, f"Excel storage is unavailable: {exc}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Run Doc.Bot on localhost with Excel storage")
    parser.add_argument("--host", default="127.0.0.1", help="Interface to bind (default: 127.0.0.1)")
    parser.add_argument("--port", type=int, default=int(os.environ.get("PORT", "8000")), help="Port (default: 8000)")
    args = parser.parse_args()
    STATIC_DIR.mkdir(parents=True, exist_ok=True)
    with LOCK:
        workbook = open_or_create_workbook()
        save_workbook(workbook)
    server = ThreadingHTTPServer((args.host, args.port), DocBotHandler)
    print(f"Doc.Bot is running at http://{args.host}:{args.port}")
    print(f"Excel database: {DB_PATH}")
    print("Keep this terminal open. Press Ctrl+C to stop.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopping Doc.Bot...")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
